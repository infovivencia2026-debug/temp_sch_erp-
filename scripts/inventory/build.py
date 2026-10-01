"""Build the inventory workbook from data.json (analyze.py) and findings.json
(the hand-written audit notes, checks and changes). See analyze.py."""
import json, datetime, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import pathlib
HERE = pathlib.Path(__file__).resolve().parent
D = json.load(open(HERE/'data.json')); extra = json.load(open(HERE/'findings.json'))
wb = Workbook()
F_ = Font(name='Arial', size=10); B_ = Font(name='Arial', size=10, bold=True); H1 = Font(name='Arial', size=14, bold=True)
HEAD = PatternFill('solid', start_color='1F2937'); HF = Font(name='Arial', size=10, bold=True, color='FFFFFF')
TONE = {'Good': 'DCFCE7', 'Check': 'FEF3C7', 'Gap': 'FEE2E2', 'Planned': 'E5E7EB', 'Fixed today': 'DBEAFE', 'Added today': 'DBEAFE', 'Open': 'FEE2E2', 'Pass': 'DCFCE7', 'Known': 'FEF3C7'}
def sheet(name, headers, rows, widths, status_col=None):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        x = ws.cell(row=1, column=c); x.font = HF; x.fill = HEAD; x.alignment = Alignment(vertical='center', wrap_text=True)
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for x in row: x.font = F_; x.alignment = Alignment(vertical='top', wrap_text=True)
        if status_col is not None:
            v = row[status_col].value
            if v in TONE: row[status_col].fill = PatternFill('solid', start_color=TONE[v])
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions; ws.row_dimensions[1].height = 30
    return ws
J = lambda xs, n=12: '\n'.join(xs[:n]) + (f'\n… +{len(xs)-n} more' if len(xs) > n else '')

# ---- Features
rows = []
for f in D['features']:
    has = bool(f['screen'] or f['bento'])
    if not has: st, note = 'Planned', 'In the catalogue, no screen built yet.'
    elif f['missing']: st, note = 'Gap', 'Screen calls an endpoint the backend does not register: ' + ', '.join(f['missing'])
    elif f['stubs']: st, note = 'Check', f"{f['stubs']} step(s) in its backend files answer \"not ported\" (see Not ported tab)."
    elif not f['api']: st, note = 'Check', 'Screen found; its data calls are made in shared parts this scan did not follow. Open it to confirm.'
    else: st, note = 'Good', ''
    rows.append([f['role'], f['workspace'], f['section'], f['name'], st, f['tier'], f['priority'], f['scope'], f['key'], f['screen'], f['bento'], len(f['api']), J(f['api'], 8), J(f['route_files'], 6), len(f['tables']), J(f['tables'], 10), J(f['worker_tests'] + f['web_tests'], 6), note, f['summary']])
sheet('Features', ['Role', 'Workspace', 'Section', 'Feature', 'Status', 'Tier', 'Priority', 'Data scope', 'Catalogue key', 'Frontend screen', 'Home board (bento)', 'API calls', 'API endpoints used', 'Backend files', 'Tables', 'Database tables touched', 'Tests that exercise it', 'Note', 'What the user sees and does'],
      rows, [22, 16, 18, 30, 11, 9, 11, 22, 44, 46, 36, 8, 40, 40, 8, 34, 30, 44, 60], status_col=4)

# ---- Backend
rows = [[r['method'], '/api/v1' + r['path'], r['perm'], r['file'], r['line'], 'Yes' if r['called'] else 'No', len(r['tables'])] for r in sorted(D['routes'], key=lambda r: (r['file'], r['line']))]
sheet('Backend routes', ['Method', 'Path', 'Permission needed', 'File', 'Line', 'Called by the web app', 'Tables its file touches'], rows, [9, 60, 34, 46, 7, 14, 12], status_col=None)

# ---- Database
rows = []
for t in D['tables']:
    st = 'Good' if t['used_by'] else 'Check'
    rows.append(['CONTROL (platform)' if t['scope'] == 'control' else 'School (one per school)', t['name'], st, t['cols'], t['idx'], t['fks'], t['migration'], len(t['used_by']), J(t['used_by'], 6), '' if t['used_by'] else 'No backend code reads or writes this table.'])
sheet('Database', ['Database', 'Table', 'Status', 'Columns', 'Indexes', 'Foreign keys', 'Created by', 'Backend files using it', 'Used by', 'Note'], rows, [22, 38, 9, 9, 9, 11, 34, 12, 50, 44], status_col=2)

# ---- Roles
rows = [[r['name'], r['key'], r['n']] for r in D['sysroles']]
sheet('Roles', ['Role', 'Key', 'Permissions granted'], rows, [34, 26, 14])

# ---- Not ported
rows = [[s['file'], s['line'], s['text']] for s in D['stubs']]
sheet('Not ported', ['Backend file', 'Line', 'What answers "not implemented" (HTTP 501)'], rows, [50, 7, 120])

# ---- Seller & support
sheet('Seller and support', ['Area', 'Status', 'Finding', 'What was done / what is needed', 'Where'], extra['seller'], [26, 13, 60, 70, 50], status_col=1)
# ---- Checks
sheet('Checks', ['Check', 'Result', 'Detail'], extra['checks'], [44, 10, 110], status_col=1)
# ---- Changes
sheet('Changed today', ['Area', 'Kind', 'Change', 'Where'], extra['changes'], [22, 13, 90, 56], status_col=1)

# ---- Summary (formulas over the other sheets)
ws = wb['Sheet']; ws.title = 'Summary'
ws['A1'] = 'School ERP inventory'; ws['A1'].font = H1
ws['A2'] = f"Generated {extra['date']} from the code on branch {extra['branch']} at {extra['commit']} by scripts/inventory/. Counts below are formulas over the other tabs."; ws['A2'].font = F_
lines = [
  ('Features in the catalogue', "=COUNTA(Features!D2:D10000)"),
  ('  Good', '=COUNTIF(Features!E2:E10000,"Good")'),
  ('  Check (see Note column)', '=COUNTIF(Features!E2:E10000,"Check")'),
  ('  Gap (screen calls a missing endpoint)', '=COUNTIF(Features!E2:E10000,"Gap")'),
  ('  Planned (no screen yet)', '=COUNTIF(Features!E2:E10000,"Planned")'),
  ('Roles in the catalogue', None),
  ('Backend routes', "=COUNTA('Backend routes'!B2:B10000)"),
  ('  Called by the web app', "=COUNTIF('Backend routes'!F2:F10000,\"Yes\")"),
  ('  Not called by the web app (devices, jobs, older screens)', "=COUNTIF('Backend routes'!F2:F10000,\"No\")"),
  ('Database tables', '=COUNTA(Database!B2:B10000)'),
  ('  Platform (CONTROL)', '=COUNTIF(Database!A2:A10000,"CONTROL*")'),
  ('  Per school', '=COUNTIF(Database!A2:A10000,"School*")'),
  ('  With no backend code using them', '=COUNTIF(Database!C2:C10000,"Check")'),
  ('Backend steps that answer "not implemented"', "=COUNTA('Not ported'!A2:A10000)"),
  ('Seller and support items still open', "=COUNTIF('Seller and support'!B2:B1000,\"Open\")"),
  ('Automated checks not passing', '=COUNTA(Checks!A2:A1000)-COUNTIF(Checks!B2:B1000,"Pass")-COUNTIF(Checks!B2:B1000,"Known")'),
]
# The same counts worked out here, beside each formula: a viewer that does not
# calculate (a file preview, a phone) shows the formula column empty.
def status(f):
    if not (f['screen'] or f['bento']): return 'Planned'
    if f['missing']: return 'Gap'
    if f['stubs'] or not f['api']: return 'Check'
    return 'Good'
ST = [status(f) for f in D['features']]
static = [len(ST), ST.count('Good'), ST.count('Check'), ST.count('Gap'), ST.count('Planned'), len({x['role'] for x in D['features']}),
  len(D['routes']), sum(1 for x in D['routes'] if x['called']), sum(1 for x in D['routes'] if not x['called']),
  len(D['tables']), sum(1 for t in D['tables'] if t['scope'] == 'control'), sum(1 for t in D['tables'] if t['scope'] != 'control'), sum(1 for t in D['tables'] if not t['used_by']),
  len(D['stubs']), sum(1 for x in extra['seller'] if x[1] == 'Open'), sum(1 for x in extra['checks'] if x[1] not in ('Pass', 'Known'))]
ws.cell(row=3, column=2, value='Count (live formula)').font = B_; ws.cell(row=3, column=3, value='Count when built').font = B_
r = 4
for (label, f), n in zip(lines, static):
    ws.cell(row=r, column=1, value=label).font = F_ if label.startswith('  ') else B_
    c = ws.cell(row=r, column=2, value=f if f else n); c.font = F_
    ws.cell(row=r, column=3, value=n).font = F_
    r += 1
print('summary counts:', dict(zip([l for l, _ in lines], static)))
r += 1
ws.cell(row=r, column=1, value='How to read the Status column').font = B_; r += 1
for k, v in [('Good', 'Screen exists and every endpoint it calls is registered in the backend.'), ('Check', 'Works as far as this scan can tell, with something to look at: see the Note.'), ('Gap', 'The screen calls an endpoint the backend does not have.'), ('Planned', 'Listed in the catalogue; no screen has been built.')]:
    c = ws.cell(row=r, column=1, value=k); c.font = F_; c.fill = PatternFill('solid', start_color=TONE[k]); ws.cell(row=r, column=2, value=v).font = F_; ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4); r += 1
r += 1
ws.cell(row=r, column=1, value='Limits of this inventory').font = B_; r += 1
for t in extra['limits']:
    ws.cell(row=r, column=1, value=t).font = F_; ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4); ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[r].height = 30; r += 1
ws.column_dimensions['A'].width = 58; ws.column_dimensions['B'].width = 24; ws.column_dimensions['C'].width = 60; ws.column_dimensions['D'].width = 20
wb.save(sys.argv[1])
print('saved', sys.argv[1])

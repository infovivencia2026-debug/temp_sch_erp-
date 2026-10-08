"""One workbook, a tab per role: every screen with the buttons on it.

The two CSVs answered different halves of the same question and made the
reviewer join them by hand. Here each role is a sheet, each row a screen, and
the buttons that screen carries sit in a column beside it -- so a person
reviewing the transport office reads one tab and writes in one Remarks column.
"""
import csv, io, os
from collections import OrderedDict, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

root = r"C:\Users\sony\AppData\Local\Temp\claude\c--Users-sony-projects-temp-sch-erp-\9039b9d2-69ae-4db2-8919-6e5d1f9d7553\scratchpad\wt-cf"
DOCS = os.path.join(root, 'docs')

feats = list(csv.DictReader(io.open(os.path.join(DOCS, 'FEATURE_REVIEW.csv'), encoding='utf-8-sig')))
acts = list(csv.DictReader(io.open(os.path.join(DOCS, 'ACTION_REVIEW.csv'), encoding='utf-8-sig')))

# Buttons, gathered per screen so they can sit beside the screen they are on.
by_key = defaultdict(list)
for a in acts:
    label = a['Button']
    if label.lower().startswith('export (table'):
        label = 'Export'
    by_key[a['Permission key']].append('%s: %s' % (a['Action'], label) if label.lower() != a['Action'].lower() else a['Action'])

HEAD = Font(bold=True, color='FFFFFF', size=11)
HEAD_FILL = PatternFill('solid', fgColor='1F2937')
TITLE = Font(bold=True, size=12)
WRAP = Alignment(vertical='top', wrap_text=True)
TOP = Alignment(vertical='top')
THIN = Side(style='thin', color='D1D5DB')
EDGE = Border(bottom=THIN)
NOTYET = PatternFill('solid', fgColor='FEF3C7')

COLS = [('Workspace', 16), ('Section', 18), ('Feature', 30), ('Built?', 9),
        ('What it is for', 52), ('Buttons on this screen', 34), ('Remarks', 34)]


def sheet_name(role, used):
    """Excel forbids : \\ / ? * [ ] and 31 characters."""
    name = role
    for ch in ':\\/?*[]':
        name = name.replace(ch, '-')
    name = name.strip()[:31] or 'Role'
    base, n = name, 2
    while name in used:
        tail = ' %d' % n
        name = base[:31 - len(tail)] + tail
        n += 1
    used.add(name)
    return name


wb = Workbook()
summary = wb.active
summary.title = 'Summary'

# --- the chains sheet: press here, check it landed there -------------------
from chains import CHAINS, FLOW_NOTE, LOGINS

ch = wb.create_sheet('TEST THIS FIRST - chains')
ch['A1'] = 'What reflects where'
ch['A1'].font = Font(bold=True, size=14)
ch['A2'] = ('Ordered by what actually happens in a school, not alphabetically. Press the action, then open the '
            'account in "Where to check" and confirm it arrived.')
ch['A2'].font = Font(size=10, color='6B7280')
ch['A3'] = ('Verified: "code" = the link was read in the source. "check" = it follows from how the screens are '
            'wired and wants confirming on the day.')
ch['A3'].font = Font(size=10, color='6B7280')

CH_COLS = [('Flow', 15), ('#', 4), ('1. Sign in as', 26), ('2. Go to', 30),
           ('3. Do exactly this', 46), ('4. You should see', 40),
           ('5. Then sign in as', 24), ('6. And open', 28), ('7. You should see there', 50),
           ('Verified', 9), ('Pass?', 8), ('Remarks', 28)]
for i, (label, width) in enumerate(CH_COLS, start=1):
    c = ch.cell(row=5, column=i, value=label)
    c.font = HEAD
    c.fill = HEAD_FILL
    c.alignment = TOP
    ch.column_dimensions[get_column_letter(i)].width = width

FLOW_FILL = PatternFill('solid', fgColor='EEF2FF')
r = 6
seen_flow = None
for flow, step, who, screen, action, happens, who2, screen2, expect, verified in CHAINS:
    if flow != seen_flow:
        seen_flow = flow
        c = ch.cell(row=r, column=1, value=flow + '  -  ' + FLOW_NOTE.get(flow, ''))
        c.font = Font(bold=True, size=11)
        c.alignment = TOP
        ch.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(CH_COLS))
        for i in range(1, len(CH_COLS) + 1):
            ch.cell(row=r, column=i).fill = FLOW_FILL
        r += 1
    for i, v in enumerate([flow, step, who, screen, action, happens, who2, screen2, expect, verified, '', ''], start=1):
        pass
    for i, v in enumerate([flow, step, who, screen, action, happens, who2, screen2, expect, verified, '', ''], start=1):
        cell = ch.cell(row=r, column=i, value=v)
        cell.alignment = WRAP
        cell.border = EDGE
        if i == 10 and v == 'check':
            cell.fill = NOTYET
    r += 1

ch.freeze_panes = 'A6'
ch.auto_filter.ref = 'A5:L%d' % (r - 1)
chain_rows = sum(1 for _ in CHAINS)

# --- every bell the product rings -----------------------------------------
# The chains are the big journeys. This is the rest of the product's
# cross-role wiring: each row is a thing that happens in one account and
# makes a notification appear in another. Read out of the Worker, not
# remembered, so it covers what the journeys do not.
from signals import SIGNALS

sg = wb.create_sheet('Signals - who gets told')
sg['A1'] = 'Every notification the product sends'
sg['A1'].font = Font(bold=True, size=14)
sg['A2'] = ('Read out of the server code. Each row is something done in one account that makes a bell ring in '
            'another. Do the thing, then sign in as the other person and look at the bell.')
sg['A2'].font = Font(size=10, color='6B7280')
sg['A3'] = ('Blank "What they are told" means the sentence is built at run time, so only the live app can show '
            'it -- which is worth checking reads properly.')
sg['A3'].font = Font(size=10, color='6B7280')

SG_COLS = [('Area', 20), ('What happened', 30), ('What they are told', 40),
           ('Where the notification opens', 40), ('In the code', 34), ('Pass?', 8), ('Remarks', 30)]
for i, (label, width) in enumerate(SG_COLS, start=1):
    c = sg.cell(row=5, column=i, value=label)
    c.font = HEAD
    c.fill = HEAD_FILL
    c.alignment = TOP
    sg.column_dimensions[get_column_letter(i)].width = width

sr = 6
for area, kind, land, src, title in SIGNALS:
    for i, v in enumerate([area, kind.replace('_', ' '), title, land, src, '', ''], start=1):
        cell = sg.cell(row=sr, column=i, value=v)
        cell.alignment = WRAP
        cell.border = EDGE
    sr += 1
sg.freeze_panes = 'A6'
sg.auto_filter.ref = 'A5:G%d' % (sr - 1)
signal_rows = len(SIGNALS)

roles = OrderedDict()
for f in feats:
    roles.setdefault(f['Role'], []).append(f)

used = set()
index = []
for role, rows in roles.items():
    ws = wb.create_sheet(sheet_name(role, used))
    ws['A1'] = role
    ws['A1'].font = TITLE
    ws['A2'] = ('%d menu entries, %d tabs inside them. Write anything you want changed in Remarks.'
                % (sum(1 for f in rows if not f['Feature'].lstrip().startswith('└')),
                   sum(1 for f in rows if f['Feature'].lstrip().startswith('└'))))
    ws['A2'].font = Font(size=10, color='6B7280')

    for i, (label, width) in enumerate(COLS, start=1):
        c = ws.cell(row=4, column=i, value=label)
        c.font = HEAD
        c.fill = HEAD_FILL
        c.alignment = TOP
        ws.column_dimensions[get_column_letter(i)].width = width

    r = 5
    built = 0
    for f in rows:
        if f['Built?'] == 'yes':
            built += 1
        # A sub-tab shares its parent's permission key, so it would otherwise
        # repeat the parent's buttons on every line. The buttons are listed
        # once, against the entry they belong to.
        sub = f['Feature'].lstrip().startswith('└')
        btns = [] if sub else by_key.get(f['Permission key'], [])
        vals = [f['Workspace'], f['Section'], f['Feature'], f['Built?'],
                f['What it is for'], '; '.join(btns), '']
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment = WRAP if i in (5, 6, 7) else TOP
            c.border = EDGE
            if f['Built?'] != 'yes':
                c.fill = NOTYET
        r += 1

    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = 'A4:G%d' % (r - 1)
    # Counted over menu entries only: a sub-tab is not a second screen to
    # review, and it shares its parent's key so it would double the buttons.
    entries = [f for f in rows if not f['Feature'].lstrip().startswith('└')]
    index.append((role, ws.title, len(entries), sum(1 for f in entries if f['Built?'] == 'yes'),
                  sum(len(by_key.get(f['Permission key'], [])) for f in entries)))

# --- the front sheet -------------------------------------------------------
summary['A1'] = 'WISEN ERP - testing workbook'
summary['A1'].font = Font(bold=True, size=14)
summary['A2'] = ('Start on "TEST THIS FIRST - chains": %d actions in the order a school does them, each saying '
                 'what it changes for somebody else and whose account to open to prove it. Then '
                 '"Signals - who gets told": %d notifications read out of the server, which is the rest of the '
                 'cross-role wiring the journeys do not cover.' % (chain_rows, signal_rows))
summary['A2'].font = Font(size=10, color='6B7280')
summary['A3'] = ('Then one tab per role: every screen, the tabs inside it (indented), the buttons on it, and a '
                 'Remarks column. A shaded row is a feature not built yet.')
summary['A3'].font = Font(size=10, color='6B7280')

# The logins, so a chain can actually be walked.
summary['A' + str(5)] = 'Logins for walking the chains'
summary['A5'].font = Font(bold=True, size=11)
for i, label in enumerate(['Role', 'Login', 'What it is good for'], start=1):
    c = summary.cell(row=6, column=i, value=label)
    c.font = HEAD
    c.fill = HEAD_FILL
lr = 7
for role, login, why in LOGINS:
    summary.cell(row=lr, column=1, value=role).border = EDGE
    summary.cell(row=lr, column=2, value=login).border = EDGE
    summary.cell(row=lr, column=3, value=why).border = EDGE
    lr += 1
LOGIN_END = lr + 1

summary.cell(row=LOGIN_END - 1, column=1, value='Every role').font = Font(bold=True, size=11)
for i, label in enumerate(['Role', 'Tab', 'Screens', 'Built', 'Buttons'], start=1):
    c = summary.cell(row=LOGIN_END, column=i, value=label)
    c.font = HEAD
    c.fill = HEAD_FILL
for w, col in zip((34, 34, 10, 10, 10), 'ABCDE'):
    summary.column_dimensions[col].width = w

r = LOGIN_END + 1
for role, tab, n, built, btns in index:
    summary.cell(row=r, column=1, value=role).border = EDGE
    link = summary.cell(row=r, column=2, value=tab)
    link.hyperlink = "#'%s'!A1" % tab
    link.font = Font(color='1D4ED8', underline='single')
    link.border = EDGE
    for col, v in ((3, n), (4, built), (5, btns)):
        summary.cell(row=r, column=col, value=v).border = EDGE
    r += 1
summary.cell(row=r, column=1, value='Total').font = Font(bold=True)
summary.cell(row=r, column=3, value=sum(x[2] for x in index)).font = Font(bold=True)
summary.cell(row=r, column=4, value=sum(x[3] for x in index)).font = Font(bold=True)
summary.cell(row=r, column=5, value=sum(x[4] for x in index)).font = Font(bold=True)
summary.freeze_panes = 'A%d' % (LOGIN_END + 1)

out = r"C:\Users\sony\Desktop\WISEN_ERP_review.xlsx"
try:
    wb.save(out)
except PermissionError:
    # Excel keeps a lock on an open workbook. Write beside it rather than
    # throw away the run; the person closes Excel and the next run replaces it.
    out = out.replace('.xlsx', '_new.xlsx')
    wb.save(out)
    print('(the original was open in Excel, so this went to the _new file)')
print('roles/tabs:', len(index))
print('screens:', sum(x[2] for x in index), '| buttons:', sum(x[4] for x in index))
print('->', out)
